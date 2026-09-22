"""Fix the "НТТУ" -> "НТУУ" / "NTTU" -> "NTUU" typo in emails.docx.

Covers body paragraphs, table cells (including nested tables) and headers/footers.
Matches are found in the paragraph's full text, so a word split across formatting runs
or inside a hyperlink is still replaced, and the surrounding formatting is kept.

Safe to re-run: if nothing matches, the file is not touched.
The original is copied to data/emails.backup.docx before saving.

    python fix_typos.py
    python parser.py        # rebuild data/queue.csv from the fixed document
"""

from __future__ import annotations

import os
import re
import shutil
import sys
import zipfile
from pathlib import Path

import docx
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent
DOCX_PATH = BASE_DIR / "emails.docx"
BACKUP_PATH = BASE_DIR / "data" / "emails.backup.docx"

# Most specific first; the generic ones catch any other context.
REPLACEMENTS = [
    ("ПЛ НТТУ «КПІ»", "ПЛ НТУУ «КПІ»"),
    ("НТТУ", "НТУУ"),
    ("PL NTTU", "PL NTUU"),
    ("NTTU", "NTUU"),
]
TYPO_RE = re.compile("|".join(re.escape(old) for old, _ in REPLACEMENTS))


def iter_paragraphs(container, seen: set | None = None):
    """Paragraphs of a document/cell/header, then those of its tables, recursively."""
    seen = set() if seen is None else seen
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat in row.cells
                    continue
                seen.add(id(cell._tc))
                yield from iter_paragraphs(cell, seen)


def iter_all_paragraphs(document):
    yield from iter_paragraphs(document)
    for section in document.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            # A linked header has no part of its own; touching it would create one.
            if not part.is_linked_to_previous:
                yield from iter_paragraphs(part)


def text_nodes(p_element) -> list:
    """<w:t> nodes belonging to this paragraph (including hyperlinks), not to nested ones."""
    return [
        t for t in p_element.iter(qn("w:t"))
        if next(t.iterancestors(qn("w:p")), None) is p_element
    ]


def replace_in_paragraph(p_element, old: str, new: str) -> int:
    count, search_from = 0, 0
    while True:
        nodes = text_nodes(p_element)
        full = "".join(n.text or "" for n in nodes)
        start = full.find(old, search_from)
        if start < 0:
            return count
        end = start + len(old)

        same_length = len(old) == len(new)
        pos, first = 0, True
        for node in nodes:
            text = node.text or ""
            node_start, node_end = pos, pos + len(text)
            pos = node_end
            if node_end <= start or node_start >= end:
                continue
            a, b = max(start, node_start) - node_start, min(end, node_end) - node_start
            if same_length:
                # Each character stays in the run it came from, so per-run formatting is untouched.
                piece = new[node_start + a - start : node_start + b - start]
            else:
                # Otherwise the replacement goes into the first run; the others just lose their part.
                piece = new if first else ""
            node.text = text[:a] + piece + text[b:]
            if node.text != node.text.strip():
                node.set(qn("xml:space"), "preserve")
            first = False
        count += 1
        search_from = start + len(new)


def apply_all(text: str) -> str:
    for old, new in REPLACEMENTS:
        text = text.replace(old, new)
    return text


def raw_hits(path: Path) -> int:
    """Typos left anywhere in the package's XML (independent cross-check)."""
    with zipfile.ZipFile(path) as z:
        return sum(
            len(TYPO_RE.findall(z.read(name).decode("utf-8", "replace")))
            for name in z.namelist() if name.endswith(".xml")
        )


def main() -> int:
    document = docx.Document(str(DOCX_PATH))
    paragraphs = list(iter_all_paragraphs(document))
    before = [p.text for p in paragraphs]

    counts = {old: 0 for old, _ in REPLACEMENTS}
    for paragraph in paragraphs:
        for old, new in REPLACEMENTS:
            counts[old] += replace_in_paragraph(paragraph._p, old, new)

    for (old, new), n in zip(REPLACEMENTS, counts.values()):
        print(f"{n:>4}  {old!r} -> {new!r}")
    total = sum(counts.values())
    if not total:
        print("Nothing to fix; emails.docx left unchanged.")
        return 0

    # Only the intended substitutions may have changed the text.
    after = [p.text for p in paragraphs]
    unexpected = [i for i, (b, a) in enumerate(zip(before, after)) if apply_all(b) != a]
    if unexpected:
        print(f"Aborting, unexpected text change in paragraph(s) {unexpected[:10]}; nothing saved.")
        return 1

    BACKUP_PATH.parent.mkdir(exist_ok=True)
    shutil.copyfile(DOCX_PATH, BACKUP_PATH)
    tmp = DOCX_PATH.with_name("emails.tmp.docx")
    document.save(str(tmp))
    try:
        os.replace(tmp, DOCX_PATH)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        print("Cannot overwrite emails.docx: close it in Word and run again.")
        return 1

    reopened = [p.text for p in iter_all_paragraphs(docx.Document(str(DOCX_PATH)))]
    left = raw_hits(DOCX_PATH)
    ok = reopened == after and left == 0
    print(f"Saved emails.docx ({total} replacement(s)); backup: {BACKUP_PATH.relative_to(BASE_DIR)}")
    print(f"Verify: {len(reopened)} paragraphs re-read, typos left in any XML part: {left} -> {'OK' if ok else 'CHECK!'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
